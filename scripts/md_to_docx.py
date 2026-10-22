"""Konversi Markdown sederhana → DOCX bertema Bio-Chain (judul, paragraf, daftar, tabel, kutipan, blok kode).

    py scripts/md_to_docx.py docs/REVISI_SOP_BIOCHAIN.md docs/REVISI_SOP_BIOCHAIN.docx
"""

import re
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, RGBColor

GREEN_DARK, GREEN, TEAL, BG_LIGHT, BG_ALT = "1B5E20", "2E7D32", "00695C", "F1F8E9", "F9FBE7"


def _rgb(h):
    return RGBColor.from_string(h)


def _shade(cell, hex_color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))


def _inline(par, text, size=10.5, color="212121", italic=False):
    """**tebal** dan `kode` di dalam baris."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        run = par.add_run(part[2:-2] if bold else part[1:-1] if code else part)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size - (0.5 if code else 0))
        run.font.name = "Consolas" if code else "Calibri"
        run.font.color.rgb = _rgb(TEAL if code else color)


def _table(doc, rows):
    header, body = rows[0], [r for r in rows[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _shade(c, GREEN_DARK)
        c.paragraphs[0].text = ""
        _inline(c.paragraphs[0], h.strip(), 9.5, "FFFFFF")
        for r in c.paragraphs[0].runs:
            r.bold = True
    for ri, row in enumerate(body, start=1):
        for ci in range(len(header)):
            c = t.rows[ri].cells[ci]
            if ri % 2 == 0:
                _shade(c, BG_ALT)
            _inline(c.paragraphs[0], (row[ci] if ci < len(row) else "").strip(), 9.5)
    doc.add_paragraph()


def convert(src: str, dst: str) -> None:
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            t = doc.add_table(rows=1, cols=1)
            c = t.rows[0].cells[0]
            _shade(c, "ECEFF1")
            c.paragraphs[0].text = ""
            run = c.paragraphs[0].add_run("\n".join(code))
            run.font.name, run.font.size = "Consolas", Pt(8.5)
            doc.add_paragraph()
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _table(doc, rows)
            continue
        elif ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            h = doc.add_heading(level=min(level, 3))
            h.text = ""
            run = h.add_run(ln.lstrip("#").strip())
            run.font.color.rgb = _rgb(GREEN_DARK if level == 1 else GREEN if level == 2 else TEAL)
        elif ln.startswith(">"):
            t = doc.add_table(rows=1, cols=1)
            c = t.rows[0].cells[0]
            _shade(c, BG_LIGHT)
            c.paragraphs[0].text = ""
            _inline(c.paragraphs[0], ln.lstrip("> ").strip(), 10, italic=False)
            while i + 1 < len(lines) and lines[i + 1].startswith(">"):
                i += 1
                p = c.add_paragraph()
                _inline(p, lines[i].lstrip("> ").strip(), 10)
            doc.add_paragraph()
        elif re.match(r"^\s*[-*] ", ln):
            _inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", ln))
        elif re.match(r"^\s*\d+\. ", ln):
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\. ", "", ln))
        elif ln.strip():
            _inline(doc.add_paragraph(), ln.strip())
        i += 1
    doc.save(dst)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
    print(f"OK → {sys.argv[2]}")
