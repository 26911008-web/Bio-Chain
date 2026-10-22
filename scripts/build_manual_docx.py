"""Manual Book Bio-Chain → DOCX bergambar (sampul, daftar isi Word, screenshot tersemat, callout berwarna).

    node scripts/render_manual_diagrams.mts      # sekali, bila diagram mermaid berubah
    py scripts/build_manual_docx.py              # → docs/manual/BIOCHAIN_VISUAL_MANUAL_BOOK.docx

Sumber: docs/manual/parts/*.md (urutan sama dengan scripts/build_manual_book.py).
Gambar dikompres ke JPEG (lebar maks 1600 px) agar berkas tetap ringan untuk dikirim.
"""

import io
import re
import sys
from datetime import date
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_manual_book import MANUAL, PARTS, TITLE  # noqa: E402

OUT = MANUAL / "BIOCHAIN_VISUAL_MANUAL_BOOK.docx"
GREEN_DARK, GREEN, INK, MUTED = "1B5E20", "3F6212", "1C2620", "5B6B60"
RED = "DC2626"
CALLOUT = {"info": ("EFF6FF", "2563EB"), "warn": ("FFF7ED", "EA580C")}
BADGE = re.compile(r"[①-⑳]")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|_[^_\s][^_]*_)")


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _shade(el, fill: str) -> None:
    """Latar sel tabel (tcPr) atau paragraf (pPr)."""
    pr = el._tc.get_or_add_tcPr() if hasattr(el, "_tc") else el._p.get_or_add_pPr()
    pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def _left_border(par, color: str) -> None:
    ppr = par._p.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="{color}"/></w:pBdr>'))


def _clean(text: str) -> str:
    text = re.sub(r"\\([<>|*_`#])", r"\1", text)          # escape markdown
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)  # tautan → teks saja


def add_inline(par, text: str, size: float = 10.5, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    """**tebal**, *miring*, `kode`, dan lencana ①②③ (merah tebal) di dalam satu paragraf."""
    for part in INLINE.split(_clean(text)):
        if not part:
            continue
        b, it, code = bold, italic, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True
        elif len(part) > 2 and part[0] in "*_" and part[-1] == part[0]:
            part, it = part[1:-1], True
        for chunk in re.split(f"({BADGE.pattern})", part):
            if not chunk:
                continue
            run = par.add_run(chunk)
            run.font.size = Pt(size)
            run.bold, run.italic = b, it
            if BADGE.fullmatch(chunk):
                run.bold, run.font.color.rgb, run.font.size = True, _rgb(RED), Pt(size + 2)
            elif code:
                run.font.name = "Consolas"
                run.font.size = Pt(size - 1)
                run.font.color.rgb = _rgb("7C2D12")
            else:
                run.font.color.rgb = _rgb(color)


def split_row(line: str) -> list[str]:
    """Pisah sel tabel pada '|' yang tidak di dalam `kode` dan tidak di-escape."""
    cells, buf, in_code, prev = [], "", False, ""
    for ch in line.strip().strip("|"):
        if ch == "`":
            in_code = not in_code
        if ch == "|" and not in_code and prev != "\\":
            cells.append(buf.strip()); buf = ""
        else:
            buf += ch
        prev = ch
    cells.append(buf.strip())
    return cells


def add_table(doc, rows: list[list[str]]) -> None:
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8.5 if ncol >= 5 else 9.5
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            par = cell.paragraphs[0]
            add_inline(par, r[j] if j < len(r) else "", size=size, bold=(i == 0),
                       color="FFFFFF" if i == 0 else INK)
            _shade(cell, GREEN if i == 0 else ("F7FAF5" if i % 2 else "FFFFFF"))
    # baris judul diulang di halaman berikutnya
    tr_pr = t.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph()


def picture_bytes(im: Image.Image) -> io.BytesIO:
    im = im.convert("RGB")
    if im.width > 1600:
        im = im.resize((1600, round(im.height * 1600 / im.width)), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=80, optimize=True)
    buf.seek(0)
    return buf


def add_picture(doc, path: Path, caption: str) -> None:
    if not path.exists():
        add_inline(doc.add_paragraph(), f"[gambar tidak ditemukan: {path.name}]", color=RED)
        return
    im = Image.open(path)
    ratio = im.height / im.width
    if ratio > 2.4:
        # Layar HP yang sangat panjang: potong dua dan susun berdampingan agar teks tetap terbaca
        half = im.height // 2
        parts = [im.crop((0, 0, im.width, half)), im.crop((0, half, im.width, im.height))]
        width = min(8.0, 21.0 / (half / im.width))
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, part in zip(t.rows[0].cells, parts):
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(picture_bytes(part), width=Cm(width))
        if caption:
            caption += " (layar panjang ditampilkan dalam dua bagian: kiri = atas, kanan = lanjutan ke bawah)"
    else:
        tall = ratio > 1.15                               # layar HP (petani/driver) → lebih sempit
        max_w, max_h = (8.0, 19.0) if tall else (16.5, 21.0)
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(picture_bytes(im), width=Cm(min(max_w, max_h / ratio)))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(cap, caption, size=8.5, color=MUTED, italic=True)


def add_callout(doc, lines: list[str]) -> None:
    text = "\n".join(lines)
    kind = "warn" if text.lstrip().lstrip("*").startswith("⚠") else "info"
    fill, border = CALLOUT[kind]
    for ln in lines:
        if not ln.strip():
            continue
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
        par = doc.add_paragraph()
        _shade(par, fill)
        _left_border(par, border)
        pf = par.paragraph_format
        pf.space_after, pf.space_before = Pt(0), Pt(0)
        pf.left_indent = Cm(0.4 + (0.5 if m else 0) + (len(m.group(1)) // 2) * 0.4 if m else 0.4)
        add_inline(par, ("• " + m.group(3)) if m else ln, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_toc(doc) -> None:
    par = doc.add_paragraph()
    run = par.add_run()
    for tag, txt in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
            if tag == "separate":
                t = OxmlElement("w:t"); t.text = "Klik kanan → Update Field untuk menampilkan daftar isi."; run._r.append(t)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)


def add_page_number(section) -> None:
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(par, f"{TITLE} · halaman ", size=8, color=MUTED)
    run = par.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.text = txt; run._r.append(it)
    run.font.size = Pt(8)


def setup(doc) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(4)
    for lvl, size, color in ((1, 20, GREEN_DARK), (2, 15, GREEN), (3, 12.5, INK), (4, 11, INK)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(size), True
        st.font.color.rgb = _rgb(color)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
    add_page_number(sec)


def cover(doc) -> None:
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "BIO-CHAIN", size=14, color=GREEN, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, TITLE, size=30, color=GREEN_DARK, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "Panduan penggunaan bergambar dengan screenshot asli beranotasi ① ② ③", size=13, color=MUTED)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "Petani / Penggilingan → Pengepul → Driver → Pabrik → sirkular abu & biochar", size=11, color=MUTED)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, f"Versi {date.today().strftime('%d-%m-%Y')} · dibangun dari docs/manual/parts/", size=9, color=MUTED)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Daftar Isi", level=1)
    add_toc(doc)


def convert(doc, md: str) -> None:
    lines = md.splitlines()
    diagrams = sorted((MANUAL / "diagrams").glob("diagram_*.png"))
    n_diag, first_h1, i = 0, True, 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("<!--"):                            # komentar editor
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1; continue
        if s.startswith("```"):
            lang = s[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i]); i += 1
            i += 1
            if lang == "mermaid" and n_diag < len(diagrams):
                add_picture(doc, diagrams[n_diag], "Diagram alur sirkular Bio-Chain"); n_diag += 1
            else:
                par = doc.add_paragraph(); _shade(par, "F3F4F6")
                run = par.add_run("\n".join(block)); run.font.name, run.font.size = "Consolas", Pt(8.5)
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|?\s*:?-{2,}", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows); continue
        if s.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(re.sub(r"^\s*> ?", "", lines[i])); i += 1
            add_callout(doc, block); continue
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if img:
            add_picture(doc, MANUAL / img.group(2).lstrip("./"), img.group(1)); i += 1; continue
        h = re.match(r"^(#{1,4})\s+(.*)", s)
        if h:
            level = len(h.group(1))
            if level == 1 and not first_h1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_h1 = first_h1 and level != 1
            hp = doc.add_heading(level=level)
            add_inline(hp, h.group(2).replace("**", ""), size={1: 20, 2: 15, 3: 12.5, 4: 11}[level],
                       color={1: GREEN_DARK, 2: GREEN}.get(level, INK), bold=True)
            i += 1; continue
        if s in ("---", "***", "___"):
            i += 1; continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", ln)
        if lm:
            depth = len(lm.group(1)) // 2
            style = "List Number" if lm.group(2)[0].isdigit() and depth == 0 else "List Bullet"
            if depth:
                style += " 2"
            par = doc.add_paragraph(style=style)
            par.paragraph_format.space_after = Pt(2)
            text = lm.group(3)
            # baris lanjutan (indentasi tanpa penanda daftar)
            while i + 1 < len(lines) and lines[i + 1].startswith("  ") and lines[i + 1].strip() \
                    and not re.match(r"^\s*([-*]|\d+\.)\s+", lines[i + 1]):
                i += 1; text += " " + lines[i].strip()
            add_inline(par, text)
            i += 1; continue
        if not s:
            i += 1; continue
        para = [s]                                          # paragraf biasa (gabung baris berurutan)
        while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^\s*(#|>|\||```|!\[|[-*] |\d+\. |<!--)", lines[i + 1]):
            i += 1; para.append(lines[i].strip())
        add_inline(doc.add_paragraph(), " ".join(para))
        i += 1


def main() -> None:
    missing = [p for p in PARTS if not (MANUAL / "parts" / p).exists()]
    if missing:
        sys.exit(f"Bagian belum ada: {', '.join(missing)}")
    md = "\n\n".join((MANUAL / "parts" / p).read_text(encoding="utf-8").strip() for p in PARTS)
    doc = docx.Document()
    setup(doc)
    cover(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    convert(doc, md)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Manual penggunaan aplikasi Bio-Chain"
    doc.save(OUT)
    print(f"DOCX: {OUT} ({OUT.stat().st_size / 1_048_576:.1f} MB)")


if __name__ == "__main__":
    main()
